from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

@dataclass
class ImportIssue:
    row_number: int
    field: str
    message: str

@dataclass
class ImportPreview:
    rows: list[dict[str, Any]] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)

def preview_staff_workbook(path: str | Path) -> ImportPreview:
    """Skeleton importer for spreadsheet migration.

    This does NOT write to the database yet.
    It provides the team with a safe 'preview then approve' starting point.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return ImportPreview(issues=[ImportIssue(0, "file", "Workbook is empty.")])

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    result = ImportPreview()

    for index, values in enumerate(rows[1:], start=2):
        record = dict(zip(headers, values))
        result.rows.append(record)

        if not any(values):
            result.issues.append(ImportIssue(index, "row", "Row is empty."))

    return result

# TODO:
# - add field mapping UI
# - validate mandatory/optional/confidential fields
# - define duplicate behaviour
# - preserve stable source IDs
# - migrate file attachments
# - add audit log and approval step
